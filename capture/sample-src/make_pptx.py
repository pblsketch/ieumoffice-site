# -*- coding: utf-8 -*-
"""북페어_발표.pptx — python-pptx 생성기 (16:9, 13.333in x 7.5in)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent.parent / "samples" / "북페어_발표.pptx"
FONT = "맑은 고딕"

INDIGO = RGBColor(0x43, 0x38, 0xCA)
INDIGO_DEEP = RGBColor(0x31, 0x2A, 0x9E)
INDIGO_SOFT = RGBColor(0xE4, 0xE2, 0xFA)
AMBER = RGBColor(0xFF, 0xB0, 0x20)
AMBER_SOFT = RGBColor(0xFF, 0xF1, 0xD2)
OFF = RGBColor(0xF7, 0xF7, 0xFB)
DARK = RGBColor(0x1E, 0x1B, 0x4B)
MUTED = RGBColor(0x5F, 0x5C, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = 13.333, 7.5


def set_run_font(run, size, bold=False, color=DARK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ko-KR")
    rpr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         line=1.15):
    """paras: list of (text, size, bold, color) or list of list-of-runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        runs = p if isinstance(p, list) else [p]
        for (t, size, bold, color) in runs:
            r = para.add_run()
            r.text = t
            set_run_font(r, size, bold, color)
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, radius=None, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def page_chrome(slide, num, section):
    """공통 상단 라벨과 하단 쪽 번호."""
    rect(slide, 0.7, 0.62, 0.42, 0.08, AMBER)
    text(slide, 1.22, 0.5, 6.0, 0.32, [(section, 12, True, INDIGO)])
    text(slide, 0.7, 6.95, 6.0, 0.3, [("이음책방 · 2026 가을 골목 북페어", 10, False, MUTED)])
    text(slide, SW - 1.7, 6.95, 1.0, 0.3, [(f"{num:02d}", 10, True, INDIGO)], align=PP_ALIGN.RIGHT)


def slide_cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, INDIGO)
    # 오른쪽 장식: 겹친 원과 책 더미
    rect(s, 8.25, 0.7, 4.7, 4.7, INDIGO_DEEP, MSO_SHAPE.OVAL)
    rect(s, 11.55, 4.75, 1.4, 1.4, AMBER, MSO_SHAPE.OVAL)
    rect(s, 8.0, 5.9, 0.55, 0.55, RGBColor(0x8B, 0x84, 0xE8), MSO_SHAPE.OVAL)
    books = [(9.55, 4.55, 2.9, 0.42, WHITE), (9.8, 4.13, 2.5, 0.42, AMBER_SOFT),
             (9.4, 3.71, 2.7, 0.42, RGBColor(0xC7, 0xC3, 0xF2)), (9.95, 3.29, 2.2, 0.42, AMBER)]
    for (bx, by, bw, bh, col) in books:
        rect(s, bx, by, bw, bh, col, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
        rect(s, bx + 0.22, by + 0.1, 0.06, bh - 0.2, INDIGO_DEEP)
    rect(s, 9.2, 4.97, 3.5, 0.06, INDIGO_DEEP)

    rect(s, 0.9, 1.45, 0.6, 0.1, AMBER)
    text(s, 0.9, 1.7, 7.0, 0.4, [("이음책방 · 개점 7주년 기획", 16, True, AMBER)])
    text(s, 0.9, 2.25, 8.0, 1.9, [("2026 가을", 40, True, WHITE), ("골목 북페어", 54, True, WHITE)],
         line=1.05)
    text(s, 0.9, 4.45, 7.6, 0.5, [("책 한 권으로 이어지는 골목의 사흘", 20, False, INDIGO_SOFT)])
    pill = rect(s, 0.9, 5.4, 5.6, 0.62, INDIGO_DEEP, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, 1.15, 5.4, 5.2, 0.62, [[("10.31 (토) – 11.2 (월)", 15, True, WHITE),
                                    ("   ·   연남 골목 일대", 15, False, INDIGO_SOFT)]],
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.9, 6.75, 6.0, 0.3, [("기획팀 김하늘 · 2026년 10월", 11, False, INDIGO_SOFT)])


def slide_overview(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, OFF)
    page_chrome(s, 2, "행사 개요")
    text(s, 0.7, 1.05, 11.5, 0.8, [("골목 전체가 하나의 책 장터가 됩니다", 32, True, DARK)])
    text(s, 0.7, 1.9, 11.0, 0.5,
         [("서점 5곳, 출판사 14곳, 공방 6곳이 사흘 동안 함께 문을 엽니다.", 16, False, MUTED)])
    cards = [
        ("일정", "10.31 – 11.2", "매일 11:00 – 20:00|우천 시 실내 부스 운영"),
        ("장소", "연남 골목 일대", "참여 공간 9곳|메인 무대: 골목 쉼터"),
        ("대상", "누구나 환영", "가족 단위 방문 환영|사전 신청 없이 입장"),
        ("주최", "이음책방 · 상인회", "후원 확정 2곳|자원봉사자 20명"),
    ]
    x0, y0, w, h, gap = 0.7, 2.8, 2.83, 3.55, 0.17
    for i, (label, big, small) in enumerate(cards):
        x = x0 + i * (w + gap)
        rect(s, x, y0, w, h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06,
             line=RGBColor(0xE2, 0xE0, 0xF4))
        rect(s, x + 0.3, y0 + 0.35, 0.62, 0.62, INDIGO_SOFT if i % 2 else AMBER_SOFT,
             MSO_SHAPE.OVAL)
        text(s, x + 0.3, y0 + 0.35, 0.62, 0.62, [(str(i + 1), 18, True, INDIGO)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, y0 + 1.2, w - 0.6, 0.35, [(label, 13, True, INDIGO)])
        text(s, x + 0.3, y0 + 1.6, w - 0.6, 0.85, [(big, 19, True, DARK)])
        rect(s, x + 0.3, y0 + 2.5, 0.5, 0.05, AMBER)
        text(s, x + 0.3, y0 + 2.7, w - 0.6, 0.75, [(ln, 12, False, MUTED) for ln in small.split("|")], line=1.25)


def slide_program(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    rect(s, 0, 0, 4.6, SH, INDIGO)
    rect(s, 0.7, 1.2, 0.6, 0.1, AMBER)
    text(s, 0.7, 1.45, 3.6, 0.4, [("PROGRAM", 14, True, AMBER)])
    text(s, 0.7, 1.95, 3.6, 2.0, [("세 갈래로", 34, True, WHITE), ("즐기는 사흘", 34, True, WHITE)],
         line=1.1)
    text(s, 0.7, 3.75, 3.5, 1.6,
         [(ln, 14, False, INDIGO_SOFT) for ln in
          ("하루 종일 골목 어디서나", "무언가가 열려 있도록", "프로그램을 배치했습니다.")], line=1.4)
    text(s, 0.7, 6.75, 3.5, 0.3, [("모든 프로그램 무료 · 체험은 현장 선착순", 10, False, INDIGO_SOFT)])

    items = [
        ("만나다", "작가와의 대화 · 책 읽는 밤",
         "동네를 쓰는 작가와 이야기를 나누고,|낭독과 어쿠스틱 공연으로 저녁을 엽니다.", AMBER),
        ("고르다", "골목 책 장터 · 중고책 교환",
         "부스 32곳에서 새 책을 고르고,|집에 있던 책을 이웃의 책과 바꿉니다.", INDIGO),
        ("만들다", "그림책 체험 · 독립출판 워크숍",
         "아이와 함께 그림책을 만들고,|나만의 작은 책을 직접 제본해 봅니다.", RGBColor(0x8B, 0x84, 0xE8)),
    ]
    y, h = 0.95, 1.75
    for i, (t, sub, desc, col) in enumerate(items):
        yy = y + i * (h + 0.22)
        rect(s, 5.3, yy, 7.35, h, OFF, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        rect(s, 5.3, yy, 0.14, h, col)
        text(s, 5.75, yy + 0.3, 1.6, 0.5, [(f"0{i + 1}", 26, True, col)])
        text(s, 7.1, yy + 0.25, 5.3, 0.45, [[(t, 20, True, DARK), ("   " + sub, 13, True, INDIGO)]])
        text(s, 7.1, yy + 0.8, 5.3, 0.75, [(ln, 13, False, MUTED) for ln in desc.split("|")], line=1.3)


def slide_stats(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, OFF)
    page_chrome(s, 4, "목표 지표")
    text(s, 0.7, 1.05, 11.5, 0.8, [("숫자로 보는 올해의 목표", 32, True, DARK)])
    text(s, 0.7, 1.9, 11.0, 0.5,
         [("지난해 봄 행사(방문 2,100명)의 두 배를 넘는 것이 목표입니다.", 16, False, MUTED)])
    stats = [
        ("4,500", "명", "누적 방문객", "사흘간 골목을 찾는 사람", INDIGO, WHITE, INDIGO_SOFT),
        ("32", "개", "참여 부스", "서점 · 출판사 · 공방", WHITE, DARK, MUTED),
        ("3,000", "권", "도서 판매", "중고책 교환 1,200권 별도", AMBER, DARK, DARK),
    ]
    x0, y0, w, h, gap = 0.7, 2.75, 3.85, 3.6, 0.2
    for i, (num, unit, label, note, fill, numc, notec) in enumerate(stats):
        x = x0 + i * (w + gap)
        rect(s, x, y0, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07,
             line=RGBColor(0xE2, 0xE0, 0xF4) if fill == WHITE else None)
        text(s, x + 0.4, y0 + 0.4, w - 0.8, 0.4, [(label, 15, True, numc)])
        text(s, x + 0.4, y0 + 1.05, w - 0.8, 1.3,
             [[(num, 60, True, numc), (" " + unit, 22, True, numc)]], anchor=MSO_ANCHOR.MIDDLE)
        rect(s, x + 0.4, y0 + 2.6, 0.6, 0.05, AMBER if fill != AMBER else DARK)
        text(s, x + 0.4, y0 + 2.8, w - 0.8, 0.4, [(note, 13, False, notec)])


def slide_timeline(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    page_chrome(s, 5, "추진 일정")
    text(s, 0.7, 1.05, 11.5, 0.8, [("행사까지 남은 길", 32, True, DARK)])
    text(s, 0.7, 1.9, 11.0, 0.5,
         [("9월부터 차근차근 준비해 11월 결과 보고까지 마칩니다.", 16, False, MUTED)])
    steps = [
        ("9월 2주", "참여 모집 마감", "부스 배치 확정"),
        ("9월 4주", "홍보 시작", "포스터 · SNS 공개"),
        ("10월 2주", "현장 준비", "봉사자 교육 · 안전 점검"),
        ("10.31 – 11.2", "북페어 개최", "사흘간 골목 운영"),
        ("11월 2주", "결과 보고", "정산 · 소책자 원고"),
    ]
    line_y = 3.85
    x0, x1 = 1.8, 11.53
    rect(s, x0, line_y - 0.03, x1 - x0, 0.06, INDIGO_SOFT)
    done = 3  # 앞의 세 단계는 완료
    rect(s, x0, line_y - 0.03, (x1 - x0) * (done - 0.5) / (len(steps) - 1), 0.06, INDIGO)
    step_w = (x1 - x0) / (len(steps) - 1)
    for i, (when, what, detail) in enumerate(steps):
        cx = x0 + i * step_w
        is_event = i == 3
        r = 0.32 if is_event else 0.22
        fill = AMBER if is_event else (INDIGO if i < done else WHITE)
        dot = rect(s, cx - r, line_y - r, 2 * r, 2 * r, fill, MSO_SHAPE.OVAL,
                   line=None if (is_event or i < done) else INDIGO)
        if not is_event and i >= done:
            dot.line.width = Pt(2.25)
        bw = 2.3
        text(s, cx - bw / 2, line_y - 1.15, bw, 0.4, [(when, 15, True, INDIGO if not is_event else
                                                         RGBColor(0xC9, 0x7A, 0x00))],
             align=PP_ALIGN.CENTER)
        card_y = line_y + 0.6
        rect(s, cx - bw / 2 + 0.05, card_y, bw - 0.1, 1.45,
             AMBER_SOFT if is_event else OFF, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        text(s, cx - bw / 2 + 0.2, card_y + 0.22, bw - 0.4, 0.4, [(what, 15, True, DARK)],
             align=PP_ALIGN.CENTER)
        text(s, cx - bw / 2 + 0.2, card_y + 0.72, bw - 0.4, 0.55, [(detail, 11.5, False, MUTED)],
             align=PP_ALIGN.CENTER)
    text(s, 0.7, 6.2, 8.0, 0.35, [[("●", 11, False, INDIGO), (" 완료    ", 11, False, MUTED),
                                    ("○", 11, False, INDIGO), (" 예정    ", 11, False, MUTED),
                                    ("●", 11, False, AMBER), (" 행사 기간", 11, False, MUTED)]])


def slide_closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, DARK)
    rect(s, 11.65, 5.2, 1.3, 1.3, INDIGO, MSO_SHAPE.OVAL)
    rect(s, 11.3, 0.8, 1.3, 1.3, AMBER, MSO_SHAPE.OVAL)
    rect(s, 1.0, 2.0, 0.6, 0.1, AMBER)
    text(s, 1.0, 2.3, 10.5, 1.0, [("골목에서 만나요", 50, True, WHITE)])
    text(s, 1.0, 3.45, 10.0, 0.6,
         [("책 한 권을 들고 오셔도, 빈손으로 오셔도 좋습니다.", 20, False, INDIGO_SOFT)])
    info = [("문의", "이음책방 기획팀 김하늘"), ("전화", "010-0000-0000"),
            ("메일", "hello@example.com")]
    for i, (k, v) in enumerate(info):
        x = 1.0 + i * 3.55
        rect(s, x, 4.7, 3.3, 0.95, RGBColor(0x2C, 0x28, 0x66), MSO_SHAPE.ROUNDED_RECTANGLE,
             radius=0.15)
        text(s, x + 0.3, 4.7, 2.9, 0.95,
             [(k, 11, True, AMBER), (v, 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE, line=1.2)


def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    slide_cover(prs)
    slide_overview(prs)
    slide_program(prs)
    slide_stats(prs)
    slide_timeline(prs)
    slide_closing(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
