# -*- coding: utf-8 -*-
"""북페어_예산표.xlsx — openpyxl 생성기."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent.parent / "samples" / "북페어_예산표.xlsx"
FONT = "맑은 고딕"
INDIGO = "4338CA"
AMBER = "FFB020"
DARK = "1E1B4B"
TINT = "F4F3FC"
NUM = "#,##0"

thin = Side(style="thin", color="D9D7F1")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(size=10.5, bold=False, color=DARK):
    return Font(name=FONT, size=size, bold=bold, color=color)


ITEMS = [
    # 항목, 구분, 수량, 단가, 비고
    ("부스 천막·테이블 대여", "시설", 32, 120000, "3일 기준"),
    ("무대 음향 장비", "시설", 3, 450000, "일 단위 대여"),
    ("골목 안내 사인물", "시설", 40, 18000, "현수막 포함"),
    ("포스터·리플릿 인쇄", "홍보", 3000, 220, "A3 포스터 300부 포함"),
    ("SNS 광고", "홍보", 1, 800000, "2주 집행"),
    ("스탬프 투어 지도", "홍보", 5000, 90, "접이식 리플릿"),
    ("작가 초청 사례비", "프로그램", 4, 500000, "대화·강연 각 2회"),
    ("공연팀 출연료", "프로그램", 2, 600000, "낭독 공연 포함"),
    ("체험 재료비", "프로그램", 300, 4500, "그림책·제본 키트"),
    ("자원봉사자 식비", "운영", 60, 9000, "20명 × 3일"),
    ("행사 보험", "운영", 1, 650000, "배상책임"),
    ("소책자 ‘골목의 책’ 제작", "기록", 500, 3800, "행사 후 발간"),
    ("완주 기념품", "운영", 800, 2500, "에코백·책갈피"),
]


def budget_sheet(ws):
    ws.title = "예산"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

    ws.merge_cells("A1:F1")
    ws["A1"] = "2026 가을 골목 북페어 · 예산 계획"
    ws["A1"].font = f(16, True, DARK)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:F2")
    ws["A2"] = "이음책방 기획팀 · 2026.10.01 기준 · 단위: 원"
    ws["A2"].font = f(9.5, False, "6B6B80")
    ws.row_dimensions[2].height = 20

    header = ["항목", "구분", "수량", "단가", "금액", "비고"]
    hr = 4
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = f(10.5, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=INDIGO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[hr].height = 24

    first = hr + 1
    for n, (name, kind, qty, price, note) in enumerate(ITEMS):
        r = first + n
        vals = [name, kind, qty, price, f"=C{r}*D{r}", note]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = f()
            c.border = BOX
            c.alignment = Alignment(vertical="center",
                                    horizontal="center" if ci == 2 else None)
            if ci in (3, 4, 5):
                c.number_format = NUM
            if n % 2 == 1:
                c.fill = PatternFill("solid", fgColor=TINT)
        ws.row_dimensions[r].height = 21
    last = first + len(ITEMS) - 1

    tr = last + 1
    ws.cell(row=tr, column=1, value="합계")
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=4)
    ws.cell(row=tr, column=5, value=f"=SUM(E{first}:E{last})")
    ws.cell(row=tr, column=6, value="예비비 별도")
    for ci in range(1, 7):
        c = ws.cell(row=tr, column=ci)
        c.font = f(11, True, DARK)
        c.fill = PatternFill("solid", fgColor="FFE7B0")
        c.border = Border(left=thin, right=thin, top=Side(style="medium", color=AMBER),
                          bottom=Side(style="medium", color=AMBER))
        c.alignment = Alignment(vertical="center",
                                horizontal="center" if ci == 1 else None)
    ws.cell(row=tr, column=5).number_format = NUM
    ws.row_dimensions[tr].height = 26

    # 구분별 소계 (차트 원본)
    sr = tr + 2
    ws.cell(row=sr, column=1, value="구분별 소계").font = f(11, True, INDIGO)
    kinds = ["시설", "홍보", "프로그램", "운영", "기록"]
    for i, h in enumerate(["구분", "금액", "비율"], start=1):
        c = ws.cell(row=sr + 1, column=i, value=h)
        c.font = f(10, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=INDIGO)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    for k, kind in enumerate(kinds):
        r = sr + 2 + k
        ws.cell(row=r, column=1, value=kind)
        ws.cell(row=r, column=2, value=f'=SUMIF($B${first}:$B${last},A{r},$E${first}:$E${last})')
        ws.cell(row=r, column=3, value=f"=B{r}/$E${tr}")
        for ci in (1, 2, 3):
            c = ws.cell(row=r, column=ci)
            c.font = f(10)
            c.border = BOX
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2).number_format = NUM
        ws.cell(row=r, column=3).number_format = "0.0%"
    kfirst, klast = sr + 2, sr + 1 + len(kinds)

    widths = {"A": 26, "B": 11, "C": 9, "D": 12, "E": 14, "F": 22}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["G"].width = 3

    ws.freeze_panes = f"A{hr + 1}"

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "구분별 예산"
    chart.y_axis.title = None
    chart.x_axis.title = None
    chart.legend = None
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.majorGridlines = None
    data = Reference(ws, min_col=2, min_row=sr + 1, max_row=klast)
    cats = Reference(ws, min_col=1, min_row=kfirst, max_row=klast)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    s = chart.series[0]
    s.graphicalProperties.solidFill = INDIGO
    s.graphicalProperties.line.solidFill = INDIGO
    chart.gapWidth = 60
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.height = 8.2
    chart.width = 13.5
    ws.add_chart(chart, "H5")
    return tr


def visitors_sheet(wb):
    ws = wb.create_sheet("방문자 예측")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "일자별 방문자 예측"
    ws["A1"].font = f(16, True, DARK)
    ws.row_dimensions[1].height = 34
    ws["A2"] = "지난해 봄 행사 실적(2,100명)과 사전 관심 등록 수를 바탕으로 추정"
    ws["A2"].font = f(9.5, False, "6B6B80")

    header = ["날짜", "요일", "오전(11–14시)", "오후(14–18시)", "저녁(18–20시)", "합계"]
    hr = 4
    for i, h in enumerate(header, start=1):
        c = ws.cell(row=hr, column=i, value=h)
        c.font = f(10.5, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=INDIGO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[hr].height = 24
    rows = [
        ("10월 31일", "토", 420, 980, 560),
        ("11월 1일", "일", 510, 1120, 430),
        ("11월 2일", "월", 180, 420, 380),
    ]
    for n, (d, w, a, b, c3) in enumerate(rows):
        r = hr + 1 + n
        vals = [d, w, a, b, c3, f"=SUM(C{r}:E{r})"]
        for ci, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = f()
            c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center")
            if ci >= 3:
                c.number_format = NUM
            if n % 2 == 1:
                c.fill = PatternFill("solid", fgColor=TINT)
        ws.row_dimensions[r].height = 22
    tr = hr + 1 + len(rows)
    ws.cell(row=tr, column=1, value="합계")
    ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=2)
    for col in "CDEF":
        ws[f"{col}{tr}"] = f"=SUM({col}{hr + 1}:{col}{tr - 1})"
    for ci in range(1, 7):
        c = ws.cell(row=tr, column=ci)
        c.font = f(11, True)
        c.fill = PatternFill("solid", fgColor="FFE7B0")
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
        if ci >= 3:
            c.number_format = NUM
    ws.row_dimensions[tr].height = 24

    nr = tr + 2
    notes = [
        ("목표 누적 방문", 4500),
        ("스탬프 투어 완주 목표", 800),
        ("예상 완주율", f"=B{nr + 1}/F{tr}"),
    ]
    for k, (label, v) in enumerate(notes):
        r = nr + k
        ws.cell(row=r, column=1, value=label).font = f(10, True, INDIGO)
        c = ws.cell(row=r, column=2, value=v)
        c.font = f(10, True)
        c.number_format = "0.0%" if k == 2 else NUM

    for col, w in {"A": 20, "B": 9, "C": 15, "D": 15, "E": 15, "F": 12}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{hr + 1}"

    ch = BarChart()
    ch.type = "col"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.title = "시간대별 방문자"
    ch.y_axis.numFmt = "#,##0"
    ch.y_axis.majorGridlines = None
    data = Reference(ws, min_col=3, max_col=5, min_row=hr, max_row=hr + len(rows))
    cats = Reference(ws, min_col=1, min_row=hr + 1, max_row=hr + len(rows))
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for s, col in zip(ch.series, ["C7C3F2", INDIGO, AMBER]):
        s.graphicalProperties.solidFill = col
        s.graphicalProperties.line.solidFill = col
    ch.gapWidth = 70
    ch.height = 8
    ch.width = 14
    ws.add_chart(ch, "H5")


def main():
    wb = Workbook()
    budget_sheet(wb.active)
    visitors_sheet(wb)
    wb.active = 0
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
