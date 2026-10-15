# -*- coding: utf-8 -*-
"""가을북페어_기획서.docx — python-docx 생성기."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "samples" / "가을북페어_기획서.docx"
FONT = "맑은 고딕"
INDIGO = RGBColor(0x43, 0x38, 0xCA)
DARK = RGBColor(0x1E, 0x1B, 0x4B)
GREY = RGBColor(0x5B, 0x5B, 0x6E)


def set_font(run_or_style, size=None, bold=None, color=None):
    font = run_or_style.font
    font.name = FONT
    rpr = run_or_style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = color


def shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def cell_margins(table, top=90, bottom=90, left=140, right=140):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def borders(table, color="D9D7F1"):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        b.append(el)
    tblpr.append(b)


def para_border_bottom(paragraph, color="4338CA", size=12):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def write_cell(cell, text, bold=False, color=DARK, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def styled_table(doc, header, rows, widths_cm, first_col_bold=True):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    cell_margins(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, "4338CA")
        write_cell(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=10,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            c = t.rows[ri].cells[ci]
            if ri % 2 == 0:
                shade(c, "F4F3FC")
            write_cell(c, val, bold=(first_col_bold and ci == 0),
                       color=INDIGO if (first_col_bold and ci == 0) else DARK)
    for row in t.rows:
        for ci, w in enumerate(widths_cm):
            row.cells[ci].width = Cm(w)
    return t


def heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=15, bold=True, color=INDIGO)
    para_border_bottom(p, "C7C3F2", 6)
    return p


def subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, size=12, bold=True, color=DARK)
    return p


def body(doc, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.45
    r = p.add_run(text)
    set_font(r, size=10.5, color=DARK)
    return p


def bullets(doc, items):
    for label, rest in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.35
        if label:
            r = p.add_run(label + "  ")
            set_font(r, size=10.5, bold=True, color=INDIGO)
        r = p.add_run(rest)
        set_font(r, size=10.5, color=DARK)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.3)

    normal = doc.styles["Normal"]
    set_font(normal, size=10.5, color=DARK)
    lb = doc.styles["List Bullet"]
    set_font(lb, size=10.5)

    # 머리말·꼬리말
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.add_run("이음책방 · 2026 가을 골목 북페어 기획서"), size=8.5, color=GREY)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run("이음책방 기획팀 · hello@example.com"), size=8.5, color=GREY)

    # 표지 블록
    tag = doc.add_paragraph()
    tag.paragraph_format.space_after = Pt(4)
    set_font(tag.add_run("행사 기획서  |  초안 v1.2"), size=10, bold=True,
             color=RGBColor(0xC9, 0x7A, 0x00))

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    set_font(title.add_run("2026 가을 골목 북페어"), size=26, bold=True, color=DARK)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    set_font(sub.add_run("책 한 권으로 이어지는 골목의 사흘 — 10월 31일(토) ~ 11월 2일(월)"),
             size=12.5, color=INDIGO)
    para_border_bottom(sub, "4338CA", 18)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(6)
    meta.paragraph_format.space_after = Pt(4)
    for k, v in (("작성", "김하늘 (이음책방 기획팀)"), ("작성일", "2026년 10월 1일"),
                 ("검토", "정은솔 대표")):
        set_font(meta.add_run(f"{k} "), size=9.5, bold=True, color=GREY)
        set_font(meta.add_run(f"{v}     "), size=9.5, color=GREY)

    # 1. 행사 개요
    heading(doc, "1. 행사 개요")
    body(doc, "이음책방은 개점 7주년을 맞아 책방이 자리한 연남 골목 전체를 하나의 책 장터로 "
              "엽니다. 동네 독립서점 다섯 곳과 소규모 출판사, 카페와 공방이 함께 참여해 "
              "사흘 동안 책을 사고팔고, 작가를 만나고, 직접 만들어 보는 자리를 마련합니다.")
    body(doc, "올해는 ‘걷는 북페어’를 콘셉트로, 방문객이 골목의 참여 공간을 돌며 스탬프를 "
              "모으고 책 이야기를 나누도록 동선을 설계했습니다. 지난해 봄 행사 방문객 "
              "2,100명의 설문 결과를 반영해 저녁 프로그램과 어린이 체험을 늘렸습니다.")

    subheading(doc, "추진 목표")
    bullets(doc, [
        ("방문객", "사흘간 누적 방문 4,500명, 스탬프 투어 완주 800명"),
        ("판매", "참여 서점·출판사 도서 판매 3,000권, 중고책 교환 1,200권"),
        ("관계", "골목 상점 12곳 이상과 협력, 지역 단골 모임 2개 신설"),
        ("기록", "행사 사진·후기를 모은 소책자 ‘골목의 책’ 500부 발간"),
    ])

    subheading(doc, "기본 정보")
    styled_table(doc, ["구분", "내용", "비고"], [
        ["일정", "2026. 10. 31.(토) ~ 11. 2.(월), 매일 11:00 ~ 20:00", "우천 시 실내 부스 운영"],
        ["장소", "이음책방 및 연남 골목 일대 (참여 공간 9곳)", "메인 무대: 골목 쉼터"],
        ["대상", "책을 좋아하는 누구나 · 가족 단위 방문객", "사전 신청 없이 입장"],
        ["주최", "이음책방 · 연남 골목 상인회", "후원 확정 2곳"],
        ["참여", "독립서점 5곳, 출판사 14곳, 공방 6곳", "부스 총 32개"],
        ["예산", "총 18,600,000원 (예비비 1,140,000원 포함)", "세부 내역은 예산표 참조"],
    ], widths_cm=[2.6, 9.0, 4.8])

    # 2. 프로그램
    heading(doc, "2. 프로그램")
    body(doc, "프로그램은 ‘만나다 · 고르다 · 만들다’ 세 갈래로 나누어 하루 종일 골목 어디서나 "
              "무언가가 열려 있도록 배치했습니다. 모든 프로그램은 무료이며, 정원이 있는 "
              "체험은 현장에서 선착순으로 접수합니다.")
    styled_table(doc, ["날짜", "시간", "프로그램", "장소"], [
        ["10/31 (토)", "11:00", "개막 · 골목 책 장터 오픈", "골목 전역"],
        ["", "15:00", "작가와의 대화 ‘동네를 쓰는 법’", "이음책방 2층"],
        ["", "19:00", "책 읽는 밤 — 낭독과 어쿠스틱 공연", "골목 쉼터 무대"],
        ["11/1 (일)", "11:00", "어린이 그림책 만들기 체험", "공방 ‘작은손’"],
        ["", "14:00", "독립출판 제작 워크숍", "카페 ‘온기’"],
        ["", "17:00", "중고책 교환 장터", "골목 쉼터"],
        ["11/2 (월)", "13:00", "북 큐레이터 추천 강연", "이음책방 2층"],
        ["", "18:00", "폐막 · 스탬프 완주 선물 증정", "골목 쉼터 무대"],
    ], widths_cm=[2.6, 1.8, 7.4, 4.6])

    subheading(doc, "운영 원칙")
    bullets(doc, [
        ("", "골목 주민의 일상을 해치지 않도록 확성 장비는 무대에서만, 20시 이전까지만 사용합니다."),
        ("", "일회용품을 줄이기 위해 장바구니 대여소와 다회용 컵 반납함을 두 곳에 운영합니다."),
        ("", "휠체어·유아차 동선을 따로 표시하고 입구마다 안내 자원봉사자를 배치합니다."),
    ])

    # 3. 추진 일정
    heading(doc, "3. 추진 일정")
    styled_table(doc, ["시기", "할 일", "담당"], [
        ["9월 2주", "참여 서점·출판사 모집 마감, 부스 배치 확정", "기획팀"],
        ["9월 4주", "포스터·안내문 제작, SNS 사전 홍보 시작", "홍보팀"],
        ["10월 2주", "자원봉사자 교육, 안전 점검, 상인회 협의", "운영팀"],
        ["10월 4주", "부스 설치 리허설, 스탬프 투어 지도 배포", "전원"],
        ["11월 2주", "결과 보고, 정산, 소책자 원고 수집", "기획팀"],
    ], widths_cm=[2.6, 10.2, 3.6])

    heading(doc, "4. 기대 효과")
    body(doc, "북페어는 책방 한 곳의 행사가 아니라 골목 전체가 함께 여는 축제입니다. 방문객은 "
              "평소 지나치던 작은 서점과 공방을 알게 되고, 상점들은 새로운 단골을 만납니다. "
              "행사가 끝난 뒤에도 이어질 독서 모임과 협력 관계를 남기는 것이 이번 기획의 "
              "가장 큰 목표입니다.")
    body(doc, "문의: 이음책방 기획팀 김하늘 · 010-0000-0000 · hello@example.com", after=0)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
